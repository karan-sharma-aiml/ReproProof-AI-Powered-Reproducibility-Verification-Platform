from __future__ import annotations

from io import BytesIO

from pptx import Presentation
from pptx.util import Inches, Pt

from .models import EnterpriseOverview


class JudgePresentationExporter:
    """Creates a compact, judge-ready presentation from a validated overview."""

    def export(self, overview: EnterpriseOverview) -> bytes:
        presentation = Presentation()
        self._title(presentation, overview)
        self._score_slide(presentation, overview)
        self._recommendation_slide(presentation, overview)
        output = BytesIO()
        presentation.save(output)
        return output.getvalue()

    @staticmethod
    def _title(presentation: Presentation, overview: EnterpriseOverview) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = "ReproProof AI"
        slide.placeholders[1].text = (
            f"Enterprise evaluation · {overview.verdict} · {overview.overall_score:.0f}/100"
        )

    @staticmethod
    def _score_slide(presentation: Presentation, overview: EnterpriseOverview) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Evidence-backed scorecard"
        box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8), Inches(5))
        frame = box.text_frame
        frame.clear()
        for name, score in overview.scores.items():
            paragraph = frame.add_paragraph()
            paragraph.text = f"{name.replace('_', ' ').title()}: {score:.0f}/100"
            paragraph.font.size = Pt(20)

    @staticmethod
    def _recommendation_slide(
        presentation: Presentation, overview: EnterpriseOverview
    ) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Recommended next actions"
        box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8), Inches(5))
        frame = box.text_frame
        frame.clear()
        recommendations = overview.recommendations or []
        for recommendation in recommendations[:8]:
            paragraph = frame.add_paragraph()
            paragraph.text = f"{recommendation.priority.upper()} · {recommendation.title}: {recommendation.reason}"
            paragraph.font.size = Pt(18)
